#!/usr/bin/env python3
"""Export a deliberately limited Markdown subset with genuine, per-occurrence footnotes.

Supports headings through ####, paragraphs, simple bold/italic, links, simple pipe tables,
and [^stable-key] notes. Reject unsupported syntax instead of silently losing content.
No automatic punctuation changes, fabricated author fields, or in-place overwrite.
"""
from __future__ import annotations
import argparse
from collections import Counter
from pathlib import Path
import re
import sys
from project import parse_markdown, REF

def export(src, out, profile, cn_font=None):
    from docx import Document
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    from docx.shared import Cm, Pt
    from docx.opc.part import Part
    from docx.opc.packuri import PackURI
    from docx.opc.constants import RELATIONSHIP_TYPE as RT
    from lxml import etree

    text=src.read_text(encoding='utf-8-sig')
    cn_font=cn_font or ('Songti SC' if sys.platform=='darwin' else 'SimSun')
    if out.exists() or out.is_symlink(): raise ValueError('Refusing to overwrite output')
    if not text.strip(): raise ValueError('Empty manuscript')
    if re.search(r'^\s*(```|~~~)|<!--|!\[|\$\$|^---\s*$',text,re.M):
        raise ValueError('Code fences, HTML comments, images, math blocks or YAML/rules require a richer document tool')
    body, defs=parse_markdown(text)
    refs=[m.group(1) for _,line in body for m in REF.finditer(line)]
    keys=[d['key'] for d in defs]
    if len(set(keys))!=len(keys): raise ValueError('Duplicate footnote definitions')
    if set(refs)!=set(keys): raise ValueError('Missing or unused footnote definitions')
    if any(v>1 for v in Counter(refs).values()): raise ValueError('Use distinct footnote occurrence keys for repeated sources')
    if profile=='journal' and refs: raise ValueError('journal mode cannot discard footnotes; use footnote mode')
    if any(not d['text'].strip() or REF.search(d['text']) for d in defs): raise ValueError('Empty or nested footnote')
    note_ids={key:n for n,key in enumerate(refs,1)}
    doc=Document()
    sec=doc.sections[0]; sec.page_width=Cm(21); sec.page_height=Cm(29.7)
    sec.top_margin=sec.bottom_margin=Cm(2.54); sec.left_margin=sec.right_margin=Cm(2.54)
    def font(run,size=12,bold=False,italic=False):
        run.font.name='Times New Roman'; run.font.size=Pt(size)
        run.font.bold=bold; run.font.italic=italic
        run._element.get_or_add_rPr().get_or_add_rFonts().set(qn('w:eastAsia'),cn_font)
    for sty in doc.styles:
        ppr=sty.element.find(qn('w:pPr'))
        if ppr is not None:
            for border in list(ppr.findall(qn('w:pBdr'))): ppr.remove(border)
    for name in ('Normal','Footnote Text','Heading 1','Heading 2','Heading 3','Title'):
        try: sty=doc.styles[name]
        except KeyError: continue
        sty.font.name='Times New Roman'; sty.font.size=Pt(12)
        from docx.shared import RGBColor
        sty.font.color.rgb=RGBColor(0,0,0)
        sty.element.get_or_add_rPr().get_or_add_rFonts().set(qn('w:eastAsia'),cn_font)
    normal=doc.styles['Normal'].paragraph_format
    normal.line_spacing=1.5; normal.space_after=Pt(0); normal.first_line_indent=Pt(24)
    tokens=re.compile(r'(\[\^[^\]\s]+\]|\*\*[^*]+\*\*|(?<!\*)\*[^*]+\*(?!\*)|\[[^\]\n]+\]\([^\)\n]+\))')
    def inline(p,value,size=12,allow_notes=True):
        for token in tokens.split(value):
            if not token: continue
            m=REF.fullmatch(token)
            if m:
                if not allow_notes: raise ValueError('Nested footnotes are not supported')
                r=p.add_run(); font(r,size)
                ref=OxmlElement('w:footnoteReference'); ref.set(qn('w:id'),str(note_ids[m.group(1)])); r._r.append(ref)
                r.font.superscript=True
            elif token.startswith('**') and token.endswith('**'): font(p.add_run(token[2:-2]),size,bold=True)
            elif token.startswith('*') and token.endswith('*'): font(p.add_run(token[1:-1]),size,italic=True)
            elif re.fullmatch(r'\[([^\]]+)\]\(([^\)]+)\)',token):
                m=re.fullmatch(r'\[([^\]]+)\]\(([^\)]+)\)',token)
                font(p.add_run(m.group(1)+' ('+m.group(2)+')'),size)
            else: font(p.add_run(token),size)
    lines=[line for _,line in body]
    i=0; first_title=True
    while i<len(lines):
        line=lines[i].strip()
        if not line: i+=1; continue
        hm=re.match(r'^(#{1,6})\s+(.+)$',line)
        if hm:
            depth=len(hm.group(1))
            if depth>4: raise ValueError('Headings beyond #### require a richer document tool')
            is_title=depth==1 and first_title
            p=doc.add_paragraph(style='Title' if is_title else f'Heading {max(1,depth-1)}')
            p.paragraph_format.first_line_indent=Pt(0); p.paragraph_format.keep_with_next=True
            inline(p,hm.group(2),16 if is_title else 14); first_title=False; i+=1; continue
        if line.startswith('|') and line.endswith('|'):
            rows=[]
            while i<len(lines) and lines[i].strip().startswith('|') and lines[i].strip().endswith('|'):
                row=[x.strip() for x in lines[i].strip()[1:-1].split('|')]
                if not all(re.fullmatch(r':?-{3,}:?',x) for x in row): rows.append(row)
                i+=1
            if not rows or any(len(r)!=len(rows[0]) for r in rows): raise ValueError('Malformed or complex pipe table')
            table=doc.add_table(rows=len(rows),cols=len(rows[0])); table.style='Table Grid'
            for ri,row in enumerate(rows):
                for ci,value in enumerate(row):
                    p=table.cell(ri,ci).paragraphs[0]; p.paragraph_format.first_line_indent=Pt(0); inline(p,value,10.5)
            continue
        buf=[line]; i+=1
        while i<len(lines) and lines[i].strip() and not re.match(r'^\s*(#{1,6}\s|\||[-*+]\s|>\s)',lines[i]):
            buf.append(lines[i].strip()); i+=1
        p=doc.add_paragraph(); inline(p,' '.join(buf))
    if refs:
        ns='http://schemas.openxmlformats.org/wordprocessingml/2006/main'
        root=etree.Element(qn('w:footnotes'),nsmap={'w':ns})
        for fid,typ in ((-1,'separator'),(0,'continuationSeparator')):
            node=etree.SubElement(root,qn('w:footnote')); node.set(qn('w:id'),str(fid)); node.set(qn('w:type'),typ)
            p=etree.SubElement(node,qn('w:p')); r=etree.SubElement(p,qn('w:r')); etree.SubElement(r,qn('w:'+typ))
        for d in defs:
            node=etree.SubElement(root,qn('w:footnote')); node.set(qn('w:id'),str(note_ids[d['key']]))
            p=etree.SubElement(node,qn('w:p'))
            ppr=etree.SubElement(p,qn('w:pPr')); sp=etree.SubElement(ppr,qn('w:spacing')); sp.set(qn('w:line'),'240'); sp.set(qn('w:lineRule'),'auto')
            r=etree.SubElement(p,qn('w:r')); etree.SubElement(r,qn('w:footnoteRef'))
            r=etree.SubElement(p,qn('w:r')); pr=etree.SubElement(r,qn('w:rPr'))
            rf=etree.SubElement(pr,qn('w:rFonts')); rf.set(qn('w:eastAsia'),cn_font); rf.set(qn('w:ascii'),'Times New Roman')
            sz=etree.SubElement(pr,qn('w:sz')); sz.set(qn('w:val'),'18')
            t=etree.SubElement(r,qn('w:t')); t.set('{http://www.w3.org/XML/1998/namespace}space','preserve'); t.text=' '+d['text']
        part=Part(PackURI('/word/footnotes.xml'),'application/vnd.openxmlformats-officedocument.wordprocessingml.footnotes+xml',etree.tostring(root,xml_declaration=True,encoding='UTF-8'),doc.part.package)
        doc.part.relate_to(part,RT.FOOTNOTES)
    footer=sec.footer.paragraphs[0]; footer.alignment=WD_ALIGN_PARAGRAPH.CENTER
    field=OxmlElement('w:fldSimple'); field.set(qn('w:instr'),'PAGE'); footer._p.append(field)
    actual=doc.element.xpath('//w:footnoteReference')
    if len(actual)!=len(refs): raise ValueError('A footnote was embedded in unsupported inline markup; move it outside that markup')
    out.parent.mkdir(parents=True,exist_ok=True); doc.save(out)
    return {'output':str(out.resolve()),'footnotes':len(refs),'layout':'NOT_ASSESSED','profile':profile}

def main():
    ap=argparse.ArgumentParser(description=__doc__); ap.add_argument('src',type=Path); ap.add_argument('-o','--output',required=True,type=Path); ap.add_argument('--profile',choices=['footnote','journal'],default='footnote'); ap.add_argument('--cn-font',default=None); a=ap.parse_args()
    try:
        import json
        print(json.dumps(export(a.src,a.output,a.profile,a.cn_font),ensure_ascii=False)); return 0
    except (OSError,ValueError,ImportError) as e:
        print('EXPORT ERROR: '+str(e),file=sys.stderr); return 2

if __name__=='__main__': sys.exit(main())
